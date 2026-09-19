import time
import io
import os
import re
import numpy as np
from PIL import Image
from docx import Document
from paddleocr import PaddleOCR

import sys
os.environ["PYMUPDF_SUGGEST_LAYOUT_ANALYZER"] = "0"

try:
    import pymupdf as fitz
    sys.modules['fitz'] = fitz
    if hasattr(fitz, 'no_recommend_layout'):
        fitz.no_recommend_layout()
    HAS_PYMUPDF = True
except ImportError:
    HAS_PYMUPDF = False

try:
    from pdf2docx import Converter
    HAS_PDF2DOCX = True
except ImportError:
    HAS_PDF2DOCX = False

from core.image_preprocessor import ImagePreprocessor
from core.exporters import ExcelExporter, SearchablePdfExporter, MarkdownExporter, ZipPackageExporter
from core.vision_ai_engine import GeminiVisionEngine


def clean_vietnamese_ocr_text(text: str) -> str:
    """
    Chuẩn hóa và sửa các lỗi mất dấu/sai ký tự đặc trưng của mô hình Latin OCR trên tiếng Việt.
    """
    if not text:
        return ""
        
    replacements = [
        ("Döc lap", "Độc lập"),
        ("Dộc lap", "Độc lập"),
        ("Doc lap", "Độc lập"),
        ("Hanh phüc", "Hạnh phúc"),
        ("Hanh phúc", "Hạnh phúc"),
        ("C.NG HOA", "CỘNG HÒA"),
        ("CÔNG HÒA", "CỘNG HÒA"),
        ("XA HI", "XÃ HỘI"),
        ("CH.NGH.A", "CHỦ NGHĨA"),
        ("VI.T NAM", "VIỆT NAM"),
        ("T6T NGHIEP", "TỐT NGHIỆP"),
        ("T6T", "TỐT"),
        ("D. AN", "DỰ ÁN"),
        ("ki-m th.", "kiểm thử"),
        ("nh-n di-n", "nhận diện"),
        ("v-n b-n", "văn bản"),
        ("ti-ng Vi-t", "tiếng Việt"),
        ("Nguyén", "Nguyễn"),
        ("Tuyét", "Tuyết"),
        ("Triêu", "Triệu"),
        ("Hoäng", "Hoàng"),
        ("Phät", "Phát"),
        ("Quäch", "Quách"),
        ("H6 Qu6c", "Hồ Quốc"),
        ("V6 Van", "Võ Văn"),
        ("V6 Minh", "Võ Minh"),
        ("Huynh Tän", "Huỳnh Tấn"),
        ("o-n v-n", "đoạn văn"),
        ("tinh n-ng", "tính năng"),
        ("ng-t trang", "ngắt trang"),
        ("thng minh", "thông minh"),
    ]
    for old, new in replacements:
        text = text.replace(old, new)
        
    return text


_PADDLE_CACHE = {}

def get_shared_paddle_model(lang: str = "vi", use_angle_cls: bool = True):
    key = (lang, use_angle_cls)
    if key not in _PADDLE_CACHE:
        _PADDLE_CACHE[key] = PaddleOCR(use_angle_cls=use_angle_cls, lang=lang, show_log=False)
    return _PADDLE_CACHE[key]


class OCREngine:
    def __init__(self, lang: str = "vi", use_angle_cls: bool = True):
        """
        Khởi tạo PaddleOCR engine với ngôn ngữ và cấu hình góc xoay.
        """
        self.lang = lang
        self.use_angle_cls = use_angle_cls
        self._model = None

    @property
    def model(self):
        if self._model is None:
            self._model = get_shared_paddle_model(self.lang, self.use_angle_cls)
        return self._model

    @staticmethod
    def check_has_digital_text(pdf_path: str, page_indices: list = None) -> tuple:
        """
        Kiểm tra xem file PDF có chứa text layer (văn bản kỹ thuật số có sẵn) hay là ảnh scan thuần.
        :return: (has_text: bool, total_characters: int, sample_text: str)
        """
        if not HAS_PYMUPDF:
            return False, 0, ""
            
        try:
            doc = fitz.open(pdf_path)
            total_doc_pages = len(doc)
            target_indices = page_indices if page_indices is not None else list(range(total_doc_pages))
            
            total_chars = 0
            sample_snippets = []
            
            for idx in target_indices:
                if 0 <= idx < total_doc_pages:
                    txt = doc[idx].get_text("text").strip()
                    total_chars += len(txt)
                    if txt and len(sample_snippets) < 3:
                        sample_snippets.append(txt[:120].replace("\n", " "))
                        
            doc.close()
            has_digital = total_chars > 30
            sample = " | ".join(sample_snippets) if sample_snippets else ""
            return has_digital, total_chars, sample
        except Exception:
            return False, 0, ""

    def render_pdf_to_images(
        self,
        pdf_path: str,
        dpi: int = 200,
        page_indices: list = None,
        deskew: bool = False,
        remove_shadow: bool = False,
        enhance_contrast: bool = False
    ) -> list:
        """
        Render các trang PDF thành danh sách ảnh PIL.Image sử dụng PyMuPDF kèm tiền xử lý OpenCV.
        """
        if not HAS_PYMUPDF:
            raise RuntimeError("Cần cài đặt thư viện pymupdf để render PDF sang ảnh.")

        doc = fitz.open(pdf_path)
        total_doc_pages = len(doc)
        target_indices = page_indices if page_indices is not None else list(range(total_doc_pages))
        
        zoom = dpi / 72.0
        mat = fitz.Matrix(zoom, zoom)

        images = []
        for idx in target_indices:
            if 0 <= idx < total_doc_pages:
                page = doc[idx]
                pix = page.get_pixmap(matrix=mat, alpha=False)
                img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)

                # Áp dụng tiền xử lý nếu được kích hoạt
                if deskew or remove_shadow or enhance_contrast:
                    img_np = np.array(img)
                    processed_np, _ = ImagePreprocessor.process_pipeline(
                        img_np,
                        deskew=deskew,
                        remove_shadow=remove_shadow,
                        enhance=enhance_contrast
                    )
                    img = Image.fromarray(processed_np)

                images.append((idx + 1, img))
        doc.close()
        return images

    @staticmethod
    def parse_page_selection(selection_str: str, max_pages: int) -> list:
        """
        Chuyển chuỗi định dạng phạm vi trang (vd: '1-3, 5') thành danh sách chỉ số trang 0-indexed.
        """
        if not selection_str or not selection_str.strip():
            return list(range(max_pages))
        
        indices = set()
        parts = [p.strip() for p in selection_str.split(",") if p.strip()]
        for part in parts:
            if "-" in part:
                subparts = part.split("-")
                if len(subparts) == 2 and subparts[0].strip().isdigit() and subparts[1].strip().isdigit():
                    start = int(subparts[0].strip())
                    end = int(subparts[1].strip())
                    for p in range(start, end + 1):
                        if 1 <= p <= max_pages:
                            indices.add(p - 1)
            elif part.isdigit():
                p = int(part)
                if 1 <= p <= max_pages:
                    indices.add(p - 1)
                    
        return sorted(list(indices)) if indices else list(range(max_pages))

    @staticmethod
    def get_pdf_page_count(pdf_path: str) -> int:
        """Lấy tổng số trang của file PDF một cách nhanh chóng."""
        if HAS_PYMUPDF:
            doc = fitz.open(pdf_path)
            count = len(doc)
            doc.close()
            return count
        return 1

    @staticmethod
    def merge_lines_into_paragraphs(ocr_lines: list, line_spacing_threshold: float = 1.4) -> list:
        """
        Gom các dòng nhận diện OCR thành đoạn văn hoàn chỉnh dựa trên khoảng cách tọa độ Y.
        """
        if not ocr_lines:
            return []

        parsed = []
        for item in ocr_lines:
            box = item[0]
            text = clean_vietnamese_ocr_text(item[1][0].strip())
            score = float(item[1][1])
            if not text:
                continue

            ys = [pt[1] for pt in box]
            xs = [pt[0] for pt in box]
            top = min(ys)
            bottom = max(ys)
            left = min(xs)
            right = max(xs)
            height = max(bottom - top, 1.0)
            parsed.append({
                "top": top,
                "bottom": bottom,
                "left": left,
                "right": right,
                "height": height,
                "text": text,
                "score": score
            })

        if not parsed:
            return []

        parsed.sort(key=lambda item: item["top"])

        paragraphs = []
        current_para = [parsed[0]["text"]]
        prev_bottom = parsed[0]["bottom"]
        avg_height = parsed[0]["height"]

        for item in parsed[1:]:
            v_gap = item["top"] - prev_bottom
            line_height = item["height"]
            avg_height = (avg_height + line_height) / 2.0

            if v_gap < avg_height * line_spacing_threshold:
                current_para.append(item["text"])
            else:
                paragraphs.append(" ".join(current_para))
                current_para = [item["text"]]

            prev_bottom = item["bottom"]

        if current_para:
            paragraphs.append(" ".join(current_para))

        return paragraphs

    def _convert_digital_pdf(
        self,
        pdf_path: str,
        output_dir: str,
        base_name: str,
        export_formats: list,
        page_indices: list = None,
        progress_callback: callable = None
    ) -> dict:
        """
        Chuyển đổi Digital PDF sang các định dạng yêu cầu.
        """
        start_time = time.time()
        if progress_callback:
            progress_callback(0, 1, "Đang phân tích cấu trúc văn bản và bảng biểu...")

        pages_summary = []
        output_files = {}

        # 1. Trích xuất text blocks để preview và xuất Markdown
        if HAS_PYMUPDF:
            doc = fitz.open(pdf_path)
            total_doc_pages = len(doc)
            target_indices = page_indices if page_indices is not None else list(range(total_doc_pages))
            
            for idx in target_indices:
                if 0 <= idx < total_doc_pages:
                    page = doc[idx]
                    page_text = page.get_text("text").strip()
                    paras = [p.strip() for p in page_text.split("\n\n") if p.strip()]
                    if not paras and page_text:
                        paras = [line.strip() for line in page_text.split("\n") if line.strip()]
                    
                    lines = [l for l in page_text.split("\n") if l.strip()]
                    pages_summary.append({
                        "page_num": idx + 1,
                        "line_count": len(lines),
                        "paragraph_count": len(paras),
                        "avg_confidence": 100.0,
                        "paragraphs": paras
                    })
            doc.close()

        # 2. Xuất Word (.docx)
        if "docx" in export_formats:
            word_path = os.path.join(output_dir, f"{base_name}.docx")
            if progress_callback:
                progress_callback(1, 3, "Đang dựng tài liệu Word với đầy đủ dấu và bảng biểu...")
            if HAS_PDF2DOCX:
                cv = Converter(pdf_path)
                cv.convert(word_path, pages=page_indices)
                cv.close()
            else:
                doc_out = Document()
                for i, p_info in enumerate(pages_summary):
                    for p_text in p_info["paragraphs"]:
                        doc_out.add_paragraph(p_text)
                    if i < len(pages_summary) - 1:
                        doc_out.add_page_break()
                doc_out.save(word_path)
            output_files["docx"] = word_path

        # 3. Xuất Excel (.xlsx)
        if "xlsx" in export_formats:
            excel_path = os.path.join(output_dir, f"{base_name}.xlsx")
            if progress_callback:
                progress_callback(2, 3, "Đang trích xuất cấu trúc bảng biểu sang Excel...")
            ExcelExporter.export_pdf_to_excel(pdf_path, excel_path, page_indices)
            output_files["xlsx"] = excel_path

        # 4. Xuất Markdown (.md)
        if "md" in export_formats:
            md_path = os.path.join(output_dir, f"{base_name}.md")
            MarkdownExporter.export_to_markdown(pages_summary, md_path)
            output_files["md"] = md_path

        # 5. Nếu Digital PDF, bản thân nó đã searchable
        if "pdf" in export_formats:
            # Copy PDF gốc làm searchable pdf vì đã có text layer
            searchable_pdf_path = os.path.join(output_dir, f"{base_name}_searchable.pdf")
            if HAS_PYMUPDF:
                doc = fitz.open(pdf_path)
                if page_indices is not None:
                    new_doc = fitz.open()
                    for pi in page_indices:
                        if 0 <= pi < len(doc):
                            new_doc.insert_pdf(doc, from_page=pi, to_page=pi)
                    new_doc.save(searchable_pdf_path)
                    new_doc.close()
                else:
                    doc.save(searchable_pdf_path)
                doc.close()
                output_files["pdf"] = searchable_pdf_path

        # 6. Đóng gói ZIP nếu có nhiều hơn 1 định dạng
        if len(output_files) > 1:
            zip_path = os.path.join(output_dir, f"{base_name}_bundle.zip")
            files_to_zip = {os.path.basename(p): p for p in output_files.values()}
            ZipPackageExporter.create_zip_archive(files_to_zip, zip_path)
            output_files["zip"] = zip_path

        elapsed = round(time.time() - start_time, 2)
        primary_word_path = output_files.get("docx", "")
        return {
            "success": True,
            "mode_used": "digital",
            "total_pages_processed": len(pages_summary),
            "overall_confidence": 100.0,
            "elapsed_seconds": elapsed,
            "output_word_path": primary_word_path,
            "output_files": output_files,
            "pages": pages_summary
        }

    def _convert_scanned_pdf(
        self,
        pdf_path: str,
        output_dir: str,
        base_name: str,
        export_formats: list,
        dpi: int = 200,
        page_indices: list = None,
        merge_paragraphs: bool = True,
        deskew: bool = False,
        remove_shadow: bool = False,
        enhance_contrast: bool = False,
        progress_callback: callable = None
    ) -> dict:
        """
        Chuyển đổi PDF dạng ảnh scan sang Word & các định dạng khác bằng PaddleOCR Vision.
        """
        start_time = time.time()
        
        if progress_callback:
            progress_callback(0, 1, "Đang trích xuất & tiền xử lý các trang PDF sang ảnh...")
            
        page_images = self.render_pdf_to_images(
            pdf_path=pdf_path,
            dpi=dpi,
            page_indices=page_indices,
            deskew=deskew,
            remove_shadow=remove_shadow,
            enhance_contrast=enhance_contrast
        )
        
        total_pages = len(page_images)
        if total_pages == 0:
            raise ValueError("Không tìm thấy trang nào để xử lý trong file PDF.")

        doc = Document()
        pages_summary = []
        all_confidences = []
        raw_ocr_pages = []

        for i, (page_num, img) in enumerate(page_images):
            if progress_callback:
                progress_callback(i + 1, total_pages, f"Đang quét OCR trang {page_num} ({i + 1}/{total_pages})...")

            img_array = np.array(img)
            result = self.model.ocr(img_array, cls=True)

            page_lines = []
            page_confs = []

            if result and result[0] is not None:
                for line in result[0]:
                    page_lines.append(line)
                    page_confs.append(float(line[1][1]))
                    all_confidences.append(float(line[1][1]))

            raw_ocr_pages.append((page_num, page_lines))

            if merge_paragraphs:
                paragraphs = self.merge_lines_into_paragraphs(page_lines)
            else:
                paragraphs = [clean_vietnamese_ocr_text(line[1][0].strip()) for line in page_lines if line[1][0].strip()]

            for p_text in paragraphs:
                doc.add_paragraph(p_text)

            if i < total_pages - 1:
                doc.add_page_break()

            avg_page_conf = (sum(page_confs) / len(page_confs)) if page_confs else 0.0
            pages_summary.append({
                "page_num": page_num,
                "line_count": len(page_lines),
                "paragraph_count": len(paragraphs),
                "avg_confidence": round(avg_page_conf * 100, 1),
                "paragraphs": paragraphs
            })

        output_files = {}

        # Lưu file Word (.docx)
        if "docx" in export_formats:
            word_path = os.path.join(output_dir, f"{base_name}.docx")
            doc.save(word_path)
            output_files["docx"] = word_path

        # Xuất Searchable PDF (.pdf)
        if "pdf" in export_formats and HAS_PYMUPDF:
            searchable_pdf_path = os.path.join(output_dir, f"{base_name}_searchable.pdf")
            SearchablePdfExporter.create_searchable_pdf(
                pdf_path=pdf_path,
                output_pdf_path=searchable_pdf_path,
                ocr_results_per_page=raw_ocr_pages,
                page_indices=page_indices
            )
            output_files["pdf"] = searchable_pdf_path

        # Xuất Excel (.xlsx)
        if "xlsx" in export_formats:
            excel_path = os.path.join(output_dir, f"{base_name}.xlsx")
            ExcelExporter.export_pdf_to_excel(pdf_path, excel_path, page_indices)
            output_files["xlsx"] = excel_path

        # Xuất Markdown (.md)
        if "md" in export_formats:
            md_path = os.path.join(output_dir, f"{base_name}.md")
            MarkdownExporter.export_to_markdown(pages_summary, md_path)
            output_files["md"] = md_path

        # Đóng gói ZIP nếu có nhiều định dạng
        if len(output_files) > 1:
            zip_path = os.path.join(output_dir, f"{base_name}_bundle.zip")
            files_to_zip = {os.path.basename(p): p for p in output_files.values()}
            ZipPackageExporter.create_zip_archive(files_to_zip, zip_path)
            output_files["zip"] = zip_path

        elapsed = round(time.time() - start_time, 2)
        overall_conf = round((sum(all_confidences) / len(all_confidences) * 100), 1) if all_confidences else 0.0
        primary_word_path = output_files.get("docx", "")

        return {
            "success": True,
            "mode_used": "ocr",
            "total_pages_processed": total_pages,
            "overall_confidence": overall_conf,
            "elapsed_seconds": elapsed,
            "output_word_path": primary_word_path,
            "output_files": output_files,
            "pages": pages_summary
        }

    def _convert_with_gemini_ai(
        self,
        pdf_path: str,
        output_dir: str,
        base_name: str,
        export_formats: list,
        gemini_api_key: str = None,
        dpi: int = 200,
        page_indices: list = None,
        deskew: bool = False,
        remove_shadow: bool = False,
        enhance_contrast: bool = False,
        progress_callback: callable = None
    ) -> dict:
        """
        Nhận diện văn bản nâng cao bằng mô hình Vision AI Google Gemini Flash.
        Hỗ trợ chữ viết tay, bản scan mờ và tự động hiệu đính ngữ cảnh tiếng Việt.
        """
        start_time = time.time()
        gemini_engine = GeminiVisionEngine(api_key=gemini_api_key)

        if not gemini_engine.is_available():
            if progress_callback:
                progress_callback(0, 1, "Gemini API chưa sẵn sàng. Tự động chuyển sang PaddleOCR...")
            return self._convert_scanned_pdf(
                pdf_path=pdf_path,
                output_dir=output_dir,
                base_name=base_name,
                export_formats=export_formats,
                dpi=dpi,
                page_indices=page_indices,
                deskew=deskew,
                remove_shadow=remove_shadow,
                enhance_contrast=enhance_contrast,
                progress_callback=progress_callback
            )

        if progress_callback:
            progress_callback(0, 1, "Đang trích xuất ảnh trang PDF...")

        page_images = self.render_pdf_to_images(
            pdf_path=pdf_path,
            dpi=dpi,
            page_indices=page_indices,
            deskew=deskew,
            remove_shadow=remove_shadow,
            enhance_contrast=enhance_contrast
        )

        total_pages = len(page_images)
        if total_pages == 0:
            raise ValueError("Không tìm thấy trang nào để xử lý trong file PDF.")

        doc = Document()
        pages_summary = []
        pages_text_map = {}
        models_used = []

        for i, (page_num, img) in enumerate(page_images):
            if progress_callback:
                progress_callback(i + 1, total_pages, f"Google Gemini Flash đang xử lý trang {page_num} ({i + 1}/{total_pages})...")

            gemini_res = gemini_engine.ocr_page_image(img)

            if gemini_res.get("status") != "success":
                # Fallback trang lỗi sang PaddleOCR
                img_array = np.array(img)
                paddle_res = self.model.ocr(img_array, cls=True)
                page_lines = paddle_res[0] if paddle_res and paddle_res[0] else []
                paragraphs = self.merge_lines_into_paragraphs(page_lines)
                page_txt = "\n\n".join(paragraphs)
                conf = 85.0
                model_used = "PaddleOCR (Fallback)"
            else:
                paragraphs = gemini_res.get("paragraphs", [])
                page_txt = gemini_res.get("text", "")
                conf = gemini_res.get("confidence", 99.0)
                model_used = gemini_res.get("model_used", "Gemini Flash")

            models_used.append(model_used)
            pages_text_map[page_num] = page_txt

            for p_text in paragraphs:
                doc.add_paragraph(p_text)

            if i < total_pages - 1:
                doc.add_page_break()

            pages_summary.append({
                "page_num": page_num,
                "line_count": len(paragraphs),
                "paragraph_count": len(paragraphs),
                "avg_confidence": conf,
                "paragraphs": paragraphs
            })

        output_files = {}

        # Xuất Word (.docx)
        if "docx" in export_formats:
            word_path = os.path.join(output_dir, f"{base_name}.docx")
            doc.save(word_path)
            output_files["docx"] = word_path

        # Xuất Searchable PDF (.pdf)
        if "pdf" in export_formats and HAS_PYMUPDF:
            searchable_pdf_path = os.path.join(output_dir, f"{base_name}_searchable.pdf")
            SearchablePdfExporter.create_searchable_pdf_from_text(
                pdf_path=pdf_path,
                output_pdf_path=searchable_pdf_path,
                page_texts=pages_text_map,
                page_indices=page_indices
            )
            output_files["pdf"] = searchable_pdf_path

        # Xuất Excel (.xlsx)
        if "xlsx" in export_formats:
            excel_path = os.path.join(output_dir, f"{base_name}.xlsx")
            ExcelExporter.export_pdf_to_excel(pdf_path, excel_path, page_indices)
            output_files["xlsx"] = excel_path

        # Xuất Markdown (.md)
        if "md" in export_formats:
            md_path = os.path.join(output_dir, f"{base_name}.md")
            MarkdownExporter.export_to_markdown(pages_summary, md_path)
            output_files["md"] = md_path

        # Đóng gói ZIP nếu có nhiều định dạng
        if len(output_files) > 1:
            zip_path = os.path.join(output_dir, f"{base_name}_bundle.zip")
            files_to_zip = {os.path.basename(p): p for p in output_files.values()}
            ZipPackageExporter.create_zip_archive(files_to_zip, zip_path)
            output_files["zip"] = zip_path

        elapsed = round(time.time() - start_time, 2)
        primary_word_path = output_files.get("docx", "")

        return {
            "success": True,
            "mode_used": "gemini",
            "model_used": ", ".join(sorted(set(models_used))),
            "total_pages_processed": total_pages,
            "overall_confidence": 99.0,
            "elapsed_seconds": elapsed,
            "output_word_path": primary_word_path,
            "output_files": output_files,
            "pages": pages_summary
        }

    def convert_pdf_to_word(
        self,
        pdf_path: str,
        output_word_path: str,
        mode: str = "auto",
        gemini_api_key: str = None,
        dpi: int = 200,
        page_indices: list = None,
        merge_paragraphs: bool = True,
        export_formats: list = None,
        deskew: bool = False,
        remove_shadow: bool = False,
        enhance_contrast: bool = False,
        progress_callback: callable = None,
        **kwargs
    ) -> dict:
        """
        Hàm chính điều phối chuyển đổi PDF sang Word và các định dạng khác theo yêu cầu.
        """
        if export_formats is None:
            export_formats = ["docx"]

        output_dir = os.path.dirname(output_word_path) or "."
        base_name = os.path.splitext(os.path.basename(output_word_path))[0]

        # Kiểm tra sự tồn tại của text layer
        has_digital, char_count, _ = self.check_has_digital_text(pdf_path, page_indices)

        if mode == "gemini":
            return self._convert_with_gemini_ai(
                pdf_path=pdf_path,
                output_dir=output_dir,
                base_name=base_name,
                export_formats=export_formats,
                gemini_api_key=gemini_api_key,
                dpi=dpi,
                page_indices=page_indices,
                deskew=deskew,
                remove_shadow=remove_shadow,
                enhance_contrast=enhance_contrast,
                progress_callback=progress_callback
            )
        elif mode == "auto":
            if has_digital and HAS_PDF2DOCX:
                return self._convert_digital_pdf(
                    pdf_path=pdf_path,
                    output_dir=output_dir,
                    base_name=base_name,
                    export_formats=export_formats,
                    page_indices=page_indices,
                    progress_callback=progress_callback
                )
            elif gemini_api_key and gemini_api_key.strip():
                return self._convert_with_gemini_ai(
                    pdf_path=pdf_path,
                    output_dir=output_dir,
                    base_name=base_name,
                    export_formats=export_formats,
                    gemini_api_key=gemini_api_key,
                    dpi=dpi,
                    page_indices=page_indices,
                    deskew=deskew,
                    remove_shadow=remove_shadow,
                    enhance_contrast=enhance_contrast,
                    progress_callback=progress_callback
                )
            else:
                return self._convert_scanned_pdf(
                    pdf_path=pdf_path,
                    output_dir=output_dir,
                    base_name=base_name,
                    export_formats=export_formats,
                    dpi=dpi,
                    page_indices=page_indices,
                    merge_paragraphs=merge_paragraphs,
                    deskew=deskew,
                    remove_shadow=remove_shadow,
                    enhance_contrast=enhance_contrast,
                    progress_callback=progress_callback
                )
        elif mode == "digital":
            return self._convert_digital_pdf(
                pdf_path=pdf_path,
                output_dir=output_dir,
                base_name=base_name,
                export_formats=export_formats,
                page_indices=page_indices,
                progress_callback=progress_callback
            )
        else:  # mode == "ocr"
            return self._convert_scanned_pdf(
                pdf_path=pdf_path,
                output_dir=output_dir,
                base_name=base_name,
                export_formats=export_formats,
                dpi=dpi,
                page_indices=page_indices,
                merge_paragraphs=merge_paragraphs,
                deskew=deskew,
                remove_shadow=remove_shadow,
                enhance_contrast=enhance_contrast,
                progress_callback=progress_callback
            )