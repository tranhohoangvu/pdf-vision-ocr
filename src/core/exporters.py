import io
import os
import zipfile
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import sys
os.environ["PYMUPDF_SUGGEST_LAYOUT_ANALYZER"] = "0"

try:
    import pymupdf as fitz
    sys.modules['fitz'] = fitz
    if hasattr(fitz, 'no_recommend_layout'):
        fitz.no_recommend_layout()
except ImportError:
    pass


class ExcelExporter:
    """
    Xuất bảng biểu và dữ liệu PDF sang file Excel (.xlsx) với định dạng chuyên nghiệp.
    """

    @staticmethod
    def _apply_table_styling(worksheet, start_row: int = 1):
        """
        Định dạng bảng tính: Header màu xanh đậm, viền mỏng, chữ căn lề tự động, độ rộng co giãn.
        """
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Arial", size=10)
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1")
        )

        for row_idx, row in enumerate(worksheet.iter_rows(min_row=start_row), start=start_row):
            for cell in row:
                cell.border = thin_border
                if row_idx == start_row:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                else:
                    cell.font = regular_font
                    cell.alignment = Alignment(vertical="center")

        # Tự động căn chỉnh độ rộng cột
        for col in worksheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                max_len = max(max_len, len(val))
            worksheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

    @classmethod
    def export_pdf_to_excel(cls, pdf_path: str, output_excel_path: str, page_indices: list = None) -> bool:
        """
        Trích xuất tất cả bảng biểu từ file PDF và lưu thành file Excel (.xlsx).
        """
        doc = fitz.open(pdf_path)
        total_pages = len(doc)
        target_indices = page_indices if page_indices is not None else list(range(total_pages))

        wb = openpyxl.Workbook()
        # Xóa sheet mặc định ban đầu
        wb.remove(wb.active)

        all_tables_rows = []
        common_header = None
        has_any_table = False

        for idx in target_indices:
            if 0 <= idx < total_pages:
                page = doc[idx]
                tabs = page.find_tables()
                page_sheet_name = f"Trang_{idx + 1}"
                ws = wb.create_sheet(title=page_sheet_name)

                page_rows = []
                if tabs and len(tabs.tables) > 0:
                    has_any_table = True
                    for table in tabs:
                        extracted = table.extract()
                        for row in extracted:
                            cleaned_row = [str(c or "").strip() for c in row]
                            if any(cleaned_row):
                                page_rows.append(cleaned_row)

                if page_rows:
                    for r in page_rows:
                        ws.append(r)
                    cls._apply_table_styling(ws, start_row=1)

                    # Gom vào danh sách tổng hợp
                    if not common_header and len(page_rows) > 0:
                        common_header = page_rows[0]
                        all_tables_rows.extend(page_rows[1:])
                    elif len(page_rows) > 1:
                        all_tables_rows.extend(page_rows[1:])
                else:
                    # Nếu không tìm thấy bảng dạng kẻ khung, bóc tách text từng dòng vào cột
                    text = page.get_text("text").strip()
                    lines = [l.strip() for l in text.split("\n") if l.strip()]
                    ws.append(["Nội dung văn bản"])
                    for line in lines:
                        ws.append([line])
                    cls._apply_table_styling(ws, start_row=1)

        doc.close()

        # Nếu có nhiều hơn 1 trang bảng, tạo thêm sheet Tổng hợp (All Data)
        if common_header and all_tables_rows and len(target_indices) > 1:
            ws_summary = wb.create_sheet(title="Tổng_Hợp", index=0)
            ws_summary.append(common_header)
            for r in all_tables_rows:
                ws_summary.append(r)
            cls._apply_table_styling(ws_summary, start_row=1)

        wb.save(output_excel_path)
        return True


class SearchablePdfExporter:
    """
    Tạo file Searchable PDF (PDF có thể bôi đen copy & Ctrl+F tìm kiếm)
    bằng cách nhúng lớp text vô hình (invisible text layer) đúng tọa độ điểm ảnh.
    """

    @staticmethod
    def create_searchable_pdf(
        pdf_path: str,
        output_pdf_path: str,
        ocr_results_per_page: list,
        page_indices: list = None
    ) -> bool:
        """
        ocr_results_per_page: danh sách [(page_num, [ [box, (text, conf)], ... ]), ...]
        """
        doc = fitz.open(pdf_path)
        total_pages = len(doc)
        target_indices = page_indices if page_indices is not None else list(range(total_pages))

        # Ánh xạ kết quả OCR theo số trang (1-indexed)
        ocr_map = {res[0]: res[1] for res in ocr_results_per_page}

        for idx in target_indices:
            if 0 <= idx < total_pages:
                page_num = idx + 1
                page = doc[idx]
                lines = ocr_map.get(page_num, [])

                for line_item in lines:
                    box = line_item[0]
                    text = line_item[1][0].strip()
                    if not text:
                        continue

                    xs = [pt[0] for pt in box]
                    ys = [pt[1] for pt in box]
                    rect = fitz.Rect(min(xs), min(ys), max(xs), max(ys))
                    
                    # Chèn text vô hình: render_mode=3
                    font_size = max(rect.height * 0.75, 8.0)
                    try:
                        page.insert_textbox(
                            rect,
                            text,
                            fontsize=font_size,
                            render_mode=3  # 3 = invisible text
                        )
                    except Exception:
                        pass

        doc.save(output_pdf_path)
        doc.close()
        return True


class MarkdownExporter:
    """
    Xuất nội dung tài liệu sang định dạng Markdown (.md) có cấu trúc.
    """

    @staticmethod
    def export_to_markdown(pages_data: list, output_md_path: str) -> bool:
        """
        pages_data: list of dicts [{"page_num": i, "paragraphs": [...]}]
        """
        md_lines = ["# Tài Liệu Trích Xuất (PDF Vision OCR)\n"]
        for page in pages_data:
            p_num = page["page_num"]
            md_lines.append(f"\n## Trang {p_num}\n")
            for para in page.get("paragraphs", []):
                md_lines.append(f"{para}\n")

        with open(output_md_path, "w", encoding="utf-8") as f:
            f.write("\n".join(md_lines))
        return True


class ZipPackageExporter:
    """
    Đóng gói nhiều định dạng đầu ra thành một file .zip tiện tải về.
    """

    @staticmethod
    def create_zip_archive(files_dict: dict, output_zip_path: str) -> bool:
        """
        files_dict: {"ten_file.docx": "/duong/dan/file.docx", ...}
        """
        with zipfile.ZipFile(output_zip_path, "w", zipfile.ZIP_DEFLATED) as zipf:
            for arcname, filepath in files_dict.items():
                if os.path.exists(filepath):
                    zipf.write(filepath, arcname=arcname)
        return True

    @staticmethod
    def create_batch_archive(batch_file_maps: list, output_zip_path: str, summary_content: str = "") -> bool:
        """
        batch_file_maps: list of dicts:
          [
            {"folder_name": "tailieu_1", "files": {"tailieu_1.docx": "/path/to/tailieu_1.docx", ...}},
            ...
          ]
        """
        with zipfile.ZipFile(output_zip_path, "w", zipfile.ZIP_DEFLATED) as zipf:
            for item in batch_file_maps:
                folder = item.get("folder_name", "")
                files = item.get("files", {})
                for fname, fpath in files.items():
                    if os.path.exists(fpath):
                        arc = f"{folder}/{fname}" if folder else fname
                        zipf.write(fpath, arcname=arc)
            
            if summary_content:
                zipf.writestr("batch_summary.txt", summary_content)
        return True
